import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Paths
INPUT_MD = r"C:\Users\ssooraj\.gemini\antigravity\brain\080fd6d2-daa6-4e39-b869-e59e30c9c253\NisolAI_Presentation.md"
DOCS_DIR = r"E:\Nisol-Labs\Code\Anti-Gravity-Code\Documents"
OUTPUT_PPTX = os.path.join(DOCS_DIR, "NisolAI_Presentation.pptx")
OUTPUT_MD_COPY = os.path.join(DOCS_DIR, "NisolAI_Presentation.md")

os.makedirs(DOCS_DIR, exist_ok=True)

with open(INPUT_MD, "r", encoding="utf-8") as f:
    full_text = f.read()

# Copy MD to Documents folder
with open(OUTPUT_MD_COPY, "w", encoding="utf-8") as f:
    f.write(full_text)

# Parsing slides
slides_data = []
slide_blocks = re.split(r'\n## Slide (\d+):?\s*', full_text)

for i in range(1, len(slide_blocks), 2):
    num = slide_blocks[i]
    content = slide_blocks[i+1]
    
    if '# APPENDIX' in content:
        content = content.split('# APPENDIX')[0]
        
    lines = content.strip().split('\n')
    title = lines[0].strip()
    
    on_slide_part = ""
    notes_part = ""
    sources_part = ""
    
    m_on_slide = re.search(r'\*\*On-slide text:\*\*(.*?)(?=\*\*Speaker notes:\*\*|\Z)', content, re.DOTALL)
    if m_on_slide:
        on_slide_part = m_on_slide.group(1).strip()
        
    m_notes = re.search(r'\*\*Speaker notes:\*\*(.*?)(?=\*\*Source files:\*\*|\Z)', content, re.DOTALL)
    if m_notes:
        notes_part = m_notes.group(1).strip()
        
    m_sources = re.search(r'\*\*Source files:\*\*(.*)', content, re.DOTALL)
    if m_sources:
        sources_part = m_sources.group(1).strip()
        
    cleaned_notes = []
    for line in notes_part.split('\n'):
        line = line.strip()
        if line.startswith('>'):
            line = line.lstrip('>').strip()
        if line:
            cleaned_notes.append(line)
    speaker_notes = "\n\n".join(cleaned_notes)
    
    slides_data.append({
        "num": int(num),
        "title": title,
        "on_slide_raw": on_slide_part,
        "notes": speaker_notes,
        "sources": sources_part
    })

print(f"Parsed {len(slides_data)} slides.")

# Create Presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5) # 16:9 widescreen

# Design Tokens (Slate 900, Sky 400, Emerald 400, Amber 400)
BG_COLOR = RGBColor(15, 23, 42)        # Slate 900
CARD_BG = RGBColor(30, 41, 59)        # Slate 800
CARD_BORDER = RGBColor(51, 65, 85)    # Slate 700
ACCENT_CYAN = RGBColor(56, 189, 248)  # Sky 400
ACCENT_GREEN = RGBColor(52, 211, 153) # Emerald 400
TEXT_WHITE = RGBColor(248, 250, 252)  # Slate 50
TEXT_MUTED = RGBColor(148, 163, 184)  # Slate 400
ACCENT_YELLOW = RGBColor(251, 191, 36)# Amber 400

blank_layout = prs.slide_layouts[6]

for s_data in slides_data:
    slide = prs.slides.add_slide(blank_layout)
    
    # Add slide notes
    if s_data["notes"]:
        notes_slide = slide.notes_slide
        text_frame = notes_slide.notes_text_frame
        text_frame.text = s_data["notes"]
        if s_data["sources"]:
            text_frame.text += "\n\n[Sources]: " + s_data["sources"]
            
    # Background rectangle
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_COLOR
    bg.line.fill.background()
    
    is_title_slide = (s_data["num"] == 1)
    
    if is_title_slide:
        # Title Slide Layout
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.16), Inches(1.3), Inches(5.0), Inches(0.5))
        pill.fill.solid()
        pill.fill.fore_color.rgb = CARD_BG
        pill.line.color.rgb = ACCENT_CYAN
        tf_pill = pill.text_frame
        tf_pill.text = "ENTERPRISE AI TRANSFORMATION PLATFORM"
        p_pill = tf_pill.paragraphs[0]
        p_pill.font.size = Pt(11)
        p_pill.font.bold = True
        p_pill.font.color.rgb = ACCENT_CYAN
        p_pill.alignment = PP_ALIGN.CENTER
        
        title_box = slide.shapes.add_textbox(Inches(1.5), Inches(2.1), Inches(10.33), Inches(3.0))
        tf = title_box.text_frame
        tf.word_wrap = True
        
        p1 = tf.paragraphs[0]
        p1.text = "Nisol AI"
        p1.font.size = Pt(56)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_WHITE
        p1.alignment = PP_ALIGN.CENTER
        
        p2 = tf.add_paragraph()
        p2.text = "Your Business Growth & AI Health-Check Machine"
        p2.font.size = Pt(26)
        p2.font.bold = True
        p2.font.color.rgb = ACCENT_GREEN
        p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(14)
        
        p3 = tf.add_paragraph()
        p3.text = "Nisol Score™ & Nisol 360™ — 7-Day Architecture Sprint"
        p3.font.size = Pt(18)
        p3.font.color.rgb = ACCENT_CYAN
        p3.alignment = PP_ALIGN.CENTER
        p3.space_before = Pt(10)
        
        sub_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.0), Inches(5.3), Inches(9.33), Inches(1.3))
        sub_card.fill.solid()
        sub_card.fill.fore_color.rgb = CARD_BG
        sub_card.line.color.rgb = CARD_BORDER
        tf_sub = sub_card.text_frame
        tf_sub.word_wrap = True
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = '"Explained so simply, even a vegetable seller can see the profit."'
        p_sub.font.size = Pt(16)
        p_sub.font.bold = True
        p_sub.font.italic = True
        p_sub.font.color.rgb = TEXT_WHITE
        p_sub.alignment = PP_ALIGN.CENTER
        
        p_sub2 = tf_sub.add_paragraph()
        p_sub2.text = "No technical jargon. Just customers, orders, cash, stock, credit, and daily profit."
        p_sub2.font.size = Pt(13)
        p_sub2.font.color.rgb = TEXT_MUTED
        p_sub2.alignment = PP_ALIGN.CENTER
        p_sub2.space_before = Pt(6)
        
    else:
        # Standard Slide Layout
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.0), Inches(0.3))
        tf_tag = tag_box.text_frame
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = f"SLIDE {s_data['num']} • NISOL AI PLATFORM WALKTHROUGH"
        p_tag.font.size = Pt(10)
        p_tag.font.bold = True
        p_tag.font.color.rgb = ACCENT_CYAN
        
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.5), Inches(0.9))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        clean_title = s_data['title'].replace('#', '').strip()
        p_title.text = clean_title
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE
        
        # Left Main Content Card (On-Slide text)
        left_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.65), Inches(7.5), Inches(5.3))
        left_card.fill.solid()
        left_card.fill.fore_color.rgb = CARD_BG
        left_card.line.color.rgb = CARD_BORDER
        
        tf_content = left_card.text_frame
        tf_content.word_wrap = True
        tf_content.margin_left = Inches(0.4)
        tf_content.margin_right = Inches(0.4)
        tf_content.margin_top = Inches(0.4)
        tf_content.margin_bottom = Inches(0.4)
        
        raw_lines = [l.strip() for l in s_data["on_slide_raw"].split('\n') if l.strip()]
        
        first = True
        for line in raw_lines:
            if line.startswith('###') or line.startswith('#'):
                continue
                
            p = tf_content.paragraphs[0] if first else tf_content.add_paragraph()
            first = False
            
            if line.startswith('-') or line.startswith('*') or line.startswith('•'):
                bullet_text = line.lstrip('-*•').strip()
                p.text = "•  " + bullet_text
                p.font.size = Pt(15)
                p.font.color.rgb = TEXT_WHITE
                p.space_after = Pt(12)
            else:
                p.text = line
                p.font.size = Pt(14)
                p.font.color.rgb = TEXT_MUTED
                p.space_after = Pt(10)
                
        # Right Side "Vegetable Seller Analogy & Business Takeaway" Card
        right_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.6), Inches(1.65), Inches(3.9), Inches(5.3))
        right_card.fill.solid()
        right_card.fill.fore_color.rgb = RGBColor(24, 33, 47)
        right_card.line.color.rgb = ACCENT_GREEN
        
        tf_right = right_card.text_frame
        tf_right.word_wrap = True
        tf_right.margin_left = Inches(0.3)
        tf_right.margin_right = Inches(0.3)
        tf_right.margin_top = Inches(0.3)
        
        p_r_hdr = tf_right.paragraphs[0]
        p_r_hdr.text = "SABZIWALA ANALOGY"
        p_r_hdr.font.size = Pt(11)
        p_r_hdr.font.bold = True
        p_r_hdr.font.color.rgb = ACCENT_GREEN
        p_r_hdr.space_after = Pt(10)
        
        notes_paragraphs = [np for np in s_data["notes"].split('\n\n') if np.strip()]
        analogy_text = ""
        for np in notes_paragraphs:
            if any(k in np.lower() for k in ["sabziwala", "mandi", "tomato", "vegetable", "hotel", "khata", "delivery", "chapati", "doctor", "shop"]):
                analogy_text = np
                break
        if not analogy_text and notes_paragraphs:
            analogy_text = notes_paragraphs[0]
            
        analogy_snippet = analogy_text[:380] + ("..." if len(analogy_text) > 380 else "")
        
        p_r_body = tf_right.add_paragraph()
        p_r_body.text = f'"{analogy_snippet}"'
        p_r_body.font.size = Pt(12)
        p_r_body.font.italic = True
        p_r_body.font.color.rgb = TEXT_WHITE
        p_r_body.space_after = Pt(14)
        
        p_takeaway_hdr = tf_right.add_paragraph()
        p_takeaway_hdr.text = "BUSINESS VALUE"
        p_takeaway_hdr.font.size = Pt(10)
        p_takeaway_hdr.font.bold = True
        p_takeaway_hdr.font.color.rgb = ACCENT_YELLOW
        p_takeaway_hdr.space_after = Pt(4)
        
        p_takeaway = tf_right.add_paragraph()
        if "Business value" in s_data["notes"]:
            bv = s_data["notes"].split("Business value")[-1].replace(":", "").strip()
            bv_clean = bv.split("\n")[0][:140]
            p_takeaway.text = bv_clean
        else:
            p_takeaway.text = "Drives more customers, eliminates manual mistakes, and speeds up cash flow."
        p_takeaway.font.size = Pt(12)
        p_takeaway.font.color.rgb = TEXT_MUTED

prs.save(OUTPUT_PPTX)
print(f"Successfully generated PowerPoint presentation at: {OUTPUT_PPTX}")
